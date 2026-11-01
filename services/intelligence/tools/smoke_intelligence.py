"""Real HTTP smoke test of the Python half of V2 (no ASP.NET, no browser).

Creates a throwaway database from the real migrations, seeds an organisation with the Python synthetic generator (standing in for the C# seeder),
starts the service with uvicorn, and drives it over HTTP exactly as the backend would: XLSX + CSV upload, repeat upload, forecast, anomalies,
Excel report. Prints a PASS/FAIL line per step and exits non-zero on any failure.

    TEST_PG_ADMIN_URL=postgresql://postgres@localhost:55432/postgres python tools/smoke_intelligence.py
"""
from __future__ import annotations

import io
import os
import subprocess
import sys
import time
import uuid
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import httpx2 as httpx
import psycopg
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from opsmind_intel.synthetic import create_structure, default_patterns, generate  # noqa: E402

ADMIN = os.environ.get("TEST_PG_ADMIN_URL", "postgresql://postgres@localhost:55432/postgres")
PORT, KEY = 8765, "smoke-key-" + uuid.uuid4().hex
failures = 0


def check(name: str, ok: bool, detail: str = "") -> None:
    global failures
    failures += 0 if ok else 1
    print(f"{'PASS' if ok else 'FAIL'}  {name}{('  ' + detail) if detail else ''}")


def main() -> int:
    db = f"opsmind_smoke_{uuid.uuid4().hex[:8]}"
    admin = psycopg.connect(ADMIN, autocommit=True)
    admin.execute(f'CREATE DATABASE "{db}"')
    url = ADMIN.rpartition("/")[0] + "/" + db
    server = None
    try:
        with psycopg.connect(url, autocommit=True) as conn:
            for sql in sorted((ROOT.parents[1] / "database" / "migrations").glob("*.sql")):
                conn.execute(sql.read_text(encoding="utf-8"))
            org = conn.execute("INSERT INTO organisations (name) VALUES ('Smoke org') RETURNING id").fetchone()[0]
        from psycopg.rows import dict_row
        with psycopg.connect(url, autocommit=True, row_factory=dict_row) as conn:
            create_structure(conn, org)

        today = datetime.now(timezone.utc).date()
        start = today - timedelta(days=180)
        data = generate(start, 180, seed=7, patterns=default_patterns(today - timedelta(days=1)))
        split = start + timedelta(days=120)
        prod, down = data.production, data.downtime
        old_p, new_p = prod[prod["shift_date"] < split.isoformat()], prod[prod["shift_date"] >= split.isoformat()]
        old_d = down[down["start_time"].astype(str).str[:10] < split.isoformat()]
        new_d = down[down["start_time"].astype(str).str[:10] >= split.isoformat()]
        xlsx = io.BytesIO()
        old_p.to_excel(xlsx, index=False)

        env = {**os.environ, "DATABASE_URL": url, "INTERNAL_API_KEY": KEY, "PORT": str(PORT)}
        server = subprocess.Popen([sys.executable, "-m", "opsmind_intel"], cwd=ROOT, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        client = httpx.Client(base_url=f"http://127.0.0.1:{PORT}", timeout=120)
        for _ in range(50):
            try:
                if client.get("/health").status_code == 200:
                    break
            except httpx.TransportError:
                time.sleep(0.2)
        h = {"X-Internal-Key": KEY, "X-Organisation-Id": str(org)}

        check("health", client.get("/health").json() == {"status": "healthy"})
        check("rejects a wrong key", client.get("/forecast", headers={**h, "X-Internal-Key": "x"}).status_code == 401)

        def upload(name, content, kind, mime="text/csv"):
            return client.post("/imports", headers=h, files={"file": (name, content, mime)}, data={"kind": kind})

        r = upload("history.xlsx", xlsx.getvalue(), "production", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet").json()
        check("XLSX production upload", r["status"] == "Completed" and r["insertedRows"] == len(old_p), f"{r['insertedRows']} rows")
        r = upload("recent.csv", new_p.to_csv(index=False).encode(), "production").json()
        check("CSV production upload", r["status"] == "Completed" and r["insertedRows"] == len(new_p), f"{r['insertedRows']} rows")
        again = upload("recent-copy.csv", new_p.to_csv(index=False).encode(), "production").json()
        check("same file again is a no-op", again["alreadyImported"] is True and again["batchId"] == r["batchId"])
        for name, frame in (("history_dt.csv", old_d), ("recent_dt.csv", new_d)):
            d = upload(name, frame.to_csv(index=False).encode(), "downtime").json()
            check(f"downtime upload {name}", d["status"] == "Completed" and d["insertedRows"] == len(frame), f"{d['insertedRows']} rows")

        bad = new_p.copy().head(5)
        bad.loc[bad.index[0], "units_produced"] = -1
        bad.loc[bad.index[1], "asset_code"] = "NOPE"
        b = upload("bad.csv", bad.to_csv(index=False).encode(), "production").json()
        check("bad rows are reported, not loaded", b["rejectedRows"] == 2 and {e["code"] for e in b["errorsPreview"]} == {"out_of_range", "unknown_equipment"},
              f"rejected={b['rejectedRows']} codes={sorted({e['code'] for e in b['errorsPreview']})}")

        f = client.get("/forecast", headers=h, params={"metric": "units", "horizon": 14})
        fj = f.json()
        check("forecast", f.status_code == 200 and len(fj["forecast"]) == 14, f"model={fj.get('selectedModel')}")
        check("forecast interval ordered", all(p["lower"] <= p["point"] <= p["upper"] for p in fj["forecast"]))
        a = client.get("/anomalies", headers=h, params={"from": str(today - timedelta(days=60)), "to": str(today - timedelta(days=1))}).json()
        check("anomalies", "anomalies" in a, f"{a['totalFlags']} flags over {a['machineDaysEvaluated']} machine-days")
        rep = client.get("/reports/operations.xlsx", headers=h, params={"from": str(today - timedelta(days=60)), "to": str(today - timedelta(days=1))})
        wb = load_workbook(io.BytesIO(rep.content))
        check("excel report", rep.status_code == 200 and "Reliability" in wb.sheetnames, ", ".join(wb.sheetnames))
        with psycopg.connect(url) as conn:
            batches = conn.execute("SELECT status, count(*) FROM import_batches GROUP BY status ORDER BY 1").fetchall()
        # 2 production + 2 downtime files, plus bad.csv whose 3 valid rows loaded (the repeat upload writes no new batch)
        check("import history persisted", dict(batches).get("Completed") == 5, str(batches))
    finally:
        if server:
            server.terminate()
            server.wait(timeout=10)
        admin.execute(f'DROP DATABASE IF EXISTS "{db}" WITH (FORCE)')
        admin.close()
    print("\nALL PASSED" if failures == 0 else f"\n{failures} CHECK(S) FAILED")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
