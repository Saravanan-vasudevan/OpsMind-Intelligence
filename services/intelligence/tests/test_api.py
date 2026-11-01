import io
import uuid
from datetime import date, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from conftest import make_org
from opsmind_intel.app import app
from opsmind_intel.etl.pipeline import run_import
from opsmind_intel.synthetic import generate

KEY = "test-key"


@pytest.fixture()
def client(database_url, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", database_url)
    monkeypatch.setenv("INTERNAL_API_KEY", KEY)
    return TestClient(app, raise_server_exceptions=False)


def hdr(org, user=None, key=KEY):
    h = {"X-Internal-Key": key, "X-Organisation-Id": str(org)}
    if user:
        h["X-User-Id"] = str(user)
    return h


@pytest.fixture()
def loaded(conn, settings, small_specs):
    org = make_org(conn, specs=small_specs)
    end = date.today() - timedelta(days=1)
    data = generate(end - timedelta(days=89), 90, seed=9, specs=small_specs)
    run_import(settings.database_url, settings, org, None, "p.csv", data.production.to_csv(index=False).encode(), "production")
    run_import(settings.database_url, settings, org, None, "d.csv", data.downtime.to_csv(index=False).encode(), "downtime")
    return org


def test_health_reports_database_connectivity(client):
    assert client.get("/health").json() == {"status": "healthy"}


def test_health_is_503_when_the_database_is_unreachable(monkeypatch):
    monkeypatch.setenv("DATABASE_URL", "postgresql://postgres@localhost:1/nope?connect_timeout=1")
    r = TestClient(app).get("/health")
    assert r.status_code == 503 and r.json() == {"status": "unhealthy"}


@pytest.mark.parametrize("headers", [{}, {"X-Internal-Key": "wrong", "X-Organisation-Id": str(uuid.uuid4())}])
def test_missing_or_wrong_key_is_401(client, headers):
    r = client.get("/forecast", headers=headers)
    assert r.status_code == 401 and r.json()["code"] == "unauthorised"


def test_unconfigured_key_refuses_everything(database_url, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", database_url)
    monkeypatch.setenv("INTERNAL_API_KEY", "")
    r = TestClient(app).get("/forecast", headers={"X-Internal-Key": "", "X-Organisation-Id": str(uuid.uuid4())})
    assert r.status_code == 503


def test_bad_organisation_header_is_400(client):
    r = client.get("/forecast", headers={"X-Internal-Key": KEY, "X-Organisation-Id": "not-a-uuid"})
    assert r.status_code == 400 and r.json()["code"] == "invalid_caller"


def test_upload_then_repeat_is_idempotent(client, conn, small_specs):
    org = make_org(conn, specs=small_specs)
    frame = generate(date(2026, 3, 1), 5, seed=1, specs=small_specs).production
    data = frame.to_csv(index=False).encode()
    files = {"file": ("p.csv", data, "text/csv")}
    first = client.post("/imports", headers=hdr(org), files=files, data={"kind": "production"})
    assert first.status_code == 200, first.text
    body = first.json()
    assert body["status"] == "Completed" and body["insertedRows"] == len(frame) and body["alreadyImported"] is False
    second = client.post("/imports", headers=hdr(org), files={"file": ("again.csv", data, "text/csv")}, data={"kind": "production"})
    assert second.json()["alreadyImported"] is True and second.json()["batchId"] == body["batchId"]
    assert conn.execute("SELECT count(*) c FROM production_records p JOIN equipment e ON e.id=p.equipment_id "
                        "JOIN production_lines l ON l.id=e.production_line_id JOIN plants pl ON pl.id=l.plant_id WHERE pl.organisation_id=%s",
                        (org,)).fetchone()["c"] == len(frame)


def test_upload_with_bad_rows_reports_them(client, conn, small_specs):
    org = make_org(conn, specs=small_specs)
    df = generate(date(2026, 3, 1), 2, seed=1, specs=small_specs).production
    df.loc[0, "units_produced"] = -5
    r = client.post("/imports", headers=hdr(org), files={"file": ("p.csv", df.to_csv(index=False).encode(), "text/csv")}, data={"kind": "production"})
    body = r.json()
    assert r.status_code == 200 and body["rejectedRows"] == 1 and body["insertedRows"] == len(df) - 1
    assert body["errorsPreview"][0]["column"] == "units_produced"


def test_upload_parameter_and_file_errors(client, conn, small_specs, monkeypatch):
    org = make_org(conn, specs=small_specs)
    csv = ("p.csv", b"a,b\n1,2\n", "text/csv")
    assert client.post("/imports", headers=hdr(org), files={"file": csv}, data={"kind": "bogus"}).status_code == 400
    assert client.post("/imports", headers=hdr(org), files={"file": ("x.pdf", b"x", "application/pdf")}, data={"kind": "production"}).status_code == 415
    assert client.post("/imports", headers=hdr(org), files={"file": ("e.csv", b"", "text/csv")}, data={"kind": "production"}).status_code == 400
    assert client.post("/imports", headers=hdr(org), data={"kind": "production"}).status_code == 400   # no file
    monkeypatch.setenv("MAX_UPLOAD_BYTES", "10")
    r = client.post("/imports", headers=hdr(org), files={"file": ("big.csv", b"x" * 50, "text/csv")}, data={"kind": "production"})
    assert r.status_code == 413 and r.json()["code"] == "file_too_large"


def test_forecast_returns_a_backtested_forecast(client, loaded):
    r = client.get("/forecast", headers=hdr(loaded), params={"metric": "units", "horizon": 7})
    assert r.status_code == 200, r.text
    body = r.json()
    assert len(body["forecast"]) == 7 and any(m["selected"] for m in body["backtest"]["models"])
    assert all(f["lower"] <= f["point"] <= f["upper"] for f in body["forecast"])


@pytest.mark.parametrize("params,code", [({"metric": "nope"}, "invalid_metric"), ({"horizon": 99}, "invalid_horizon")])
def test_forecast_parameter_validation(client, loaded, params, code):
    r = client.get("/forecast", headers=hdr(loaded), params=params)
    assert r.status_code == 400 and r.json()["code"] == code


def test_forecast_with_no_data_is_422_not_a_made_up_number(client, conn, small_specs):
    org = make_org(conn, specs=small_specs)
    r = client.get("/forecast", headers=hdr(org))
    assert r.status_code == 422 and r.json()["code"] == "no_data"


def test_forecast_with_too_little_history_is_422(client, conn, settings, small_specs):
    org = make_org(conn, specs=small_specs)
    end = date.today() - timedelta(days=1)
    data = generate(end - timedelta(days=9), 10, seed=3, specs=small_specs)
    run_import(settings.database_url, settings, org, None, "p.csv", data.production.to_csv(index=False).encode(), "production")
    r = client.get("/forecast", headers=hdr(org))
    assert r.status_code == 422 and r.json()["code"] == "insufficient_history"


def test_another_organisations_plant_is_404(client, conn, loaded, small_specs):
    other = make_org(conn, specs=small_specs)
    plant = conn.execute("SELECT id FROM plants WHERE organisation_id = %s", (other,)).fetchone()["id"]
    assert client.get("/forecast", headers=hdr(loaded), params={"plantId": str(plant)}).status_code == 404
    assert client.get("/anomalies", headers=hdr(loaded), params={"plantId": str(plant)}).status_code == 404


def test_anomalies_and_parameter_errors(client, loaded):
    ok = client.get("/anomalies", headers=hdr(loaded))
    assert ok.status_code == 200 and ok.json()["method"] == "robust_z"
    assert client.get("/anomalies", headers=hdr(loaded), params={"method": "magic"}).status_code == 400
    assert client.get("/anomalies", headers=hdr(loaded), params={"from": "2026-05-01", "to": "2026-04-01"}).status_code == 400
    assert client.get("/anomalies", headers=hdr(loaded), params={"from": "garbage"}).status_code == 400


def test_excel_report_download(client, loaded):
    today = date.today()
    r = client.get("/reports/operations.xlsx", headers=hdr(loaded), params={"from": str(today - timedelta(days=30)), "to": str(today - timedelta(days=1))})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in r.headers["content-disposition"]
    assert "Summary" in load_workbook(io.BytesIO(r.content)).sheetnames
    bad = client.get("/reports/operations.xlsx", headers=hdr(loaded), params={"from": "2026-05-01", "to": "2026-04-01"})
    assert bad.status_code == 400


def test_evaluation_is_labelled_synthetic(client, loaded):
    body = client.get("/anomalies/evaluation", headers=hdr(loaded)).json()
    assert "synthetic" in body["dataset"]["kind"] and body["limitations"]
