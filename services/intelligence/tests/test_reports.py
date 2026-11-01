import io
import shutil
import subprocess
import tempfile
from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from opsmind_intel.etl.pipeline import run_import
from opsmind_intel.reports.workbook import ReportParameterError, build_report
from opsmind_intel.synthetic import generate
from conftest import make_org

START = date(2026, 3, 1)


@pytest.fixture()
def loaded(conn, settings, small_specs):
    org = make_org(conn, specs=small_specs)
    data = generate(START, 28, seed=5, specs=small_specs)
    run_import(settings.database_url, settings, org, None, "p.csv", data.production.to_csv(index=False).encode(), "production")
    run_import(settings.database_url, settings, org, None, "d.csv", data.downtime.to_csv(index=False).encode(), "downtime")
    return org


def test_workbook_has_expected_sheets_and_formula_kpis(conn, loaded):
    data = build_report(conn, loaded, START, date(2026, 3, 28), generated_at=datetime(2026, 4, 1, tzinfo=timezone.utc))
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Summary", "Daily", "Reliability", "Downtime reasons", "Anomalies", "About"]
    s = wb["Summary"]
    assert str(s["B10"].value).startswith("=SUM(Daily!B2:B")
    assert "n/a" in str(s["B11"].value)  # the zero-denominator branch exists in the formula
    assert wb["Daily"].max_row == 29   # header + 28 days
    assert "Synthetic" in s["A2"].value


def test_kpis_recalculate_to_the_database_values(conn, loaded, tmp_path):
    """Recalculate the formulas with LibreOffice and compare with an independent SQL computation."""
    soffice = shutil.which("soffice")
    if not soffice:
        pytest.skip("LibreOffice not installed")
    data = build_report(conn, loaded, START, date(2026, 3, 28))
    src = tmp_path / "r.xlsx"
    src.write_bytes(data)
    out = tmp_path / "out"
    subprocess.run([soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(src)], check=True, capture_output=True,
                   timeout=120, env={"HOME": str(tmp_path), "PATH": "/usr/bin:/bin"})
    s = load_workbook(out / "r.xlsx", data_only=True)["Summary"]
    q = conn.execute(
        """SELECT sum(p.units_produced) u, sum(p.target_units) t, sum(p.run_minutes) r, sum(sh.duration_minutes) sc
           FROM production_records p JOIN shifts sh ON sh.id = p.shift_id JOIN equipment e ON e.id = p.equipment_id
           JOIN production_lines l ON l.id = e.production_line_id JOIN plants pl ON pl.id = l.plant_id
           WHERE pl.organisation_id = %s""", (loaded,)).fetchone()
    assert s["B10"].value == q["u"]
    assert s["B11"].value == pytest.approx(q["u"] / q["t"])
    assert s["B12"].value == pytest.approx(q["r"] / q["sc"])


def test_empty_range_gives_na_not_zero_percent(conn, loaded, tmp_path):
    soffice = shutil.which("soffice")
    if not soffice:
        pytest.skip("LibreOffice not installed")
    data = build_report(conn, loaded, date(2030, 1, 1), date(2030, 1, 7))
    src = tmp_path / "e.xlsx"
    src.write_bytes(data)
    subprocess.run([soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "o"), str(src)], check=True,
                   capture_output=True, timeout=120, env={"HOME": str(tmp_path), "PATH": "/usr/bin:/bin"})
    s = load_workbook(tmp_path / "o" / "e.xlsx", data_only=True)["Summary"]
    assert s["B11"].value == "n/a" and s["B12"].value == "n/a" and s["B13"].value == "n/a" and s["B14"].value == "n/a"


def test_bad_parameters_are_rejected(conn, loaded):
    with pytest.raises(ReportParameterError):
        build_report(conn, loaded, date(2026, 3, 10), date(2026, 3, 1))
    with pytest.raises(ReportParameterError):
        build_report(conn, loaded, date(2024, 1, 1), date(2026, 3, 1))
    import uuid
    with pytest.raises(ReportParameterError):
        build_report(conn, loaded, START, date(2026, 3, 28), plant_id=uuid.uuid4())


def test_other_organisations_data_never_appears(conn, loaded, small_specs):
    other = make_org(conn, specs=small_specs)
    wb = load_workbook(io.BytesIO(build_report(conn, other, START, date(2026, 3, 28))))
    assert wb["Daily"].max_row == 1 and wb["Reliability"].max_row <= 3
