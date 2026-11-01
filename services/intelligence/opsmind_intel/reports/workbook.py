"""Excel operations report.

The Summary sheet's KPIs are live Excel formulas over the Daily sheet (so a reader can audit them), using the same
definitions as the API (docs/KPI_DEFINITIONS.md). A ratio with a zero denominator shows "n/a", never 0%.
"""
from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timezone

import psycopg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..anomaly.service import detect

MAX_DAYS = 366
HEADER_FILL = PatternFill("solid", fgColor="1F3A4D")
HEADER_FONT = Font(bold=True, color="FFFFFF")


class ReportParameterError(ValueError):
    pass


DAILY_COLUMNS = [
    ("Date", "date", "yyyy-mm-dd"), ("Units produced", "units", "#,##0"), ("Target units", "target_units", "#,##0"),
    ("Run minutes", "run_minutes", "#,##0"), ("Scheduled minutes", "scheduled_minutes", "#,##0"),
    ("Inspected units", "inspected_units", "#,##0"), ("Rejected units", "rejected_units", "#,##0"),
    ("Planned downtime (min)", "planned", "#,##0"), ("Unplanned downtime (min)", "unplanned", "#,##0"), ("Failures", "failures", "#,##0"),
]


def _header(ws, row: int, labels: list[str]) -> None:
    for i, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=label)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_report(conn: psycopg.Connection, organisation_id: uuid.UUID, start: date, end: date, plant_id: uuid.UUID | None = None,
                 include_anomalies: bool = True, generated_at: datetime | None = None) -> bytes:
    if end < start:
        raise ReportParameterError("'to' must not be before 'from'")
    if (end - start).days + 1 > MAX_DAYS:
        raise ReportParameterError(f"The report is limited to {MAX_DAYS} days")
    generated_at = generated_at or datetime.now(timezone.utc)

    plant_name = None
    if plant_id:
        row = conn.execute("SELECT name FROM plants WHERE id = %s AND organisation_id = %s", (plant_id, organisation_id)).fetchone()
        if not row:
            raise ReportParameterError("Plant not found")
        plant_name = row["name"]

    daily = conn.execute(
        """SELECT day AS date, sum(units)::bigint AS units, sum(target_units)::bigint AS target_units, sum(run_minutes)::bigint AS run_minutes,
                  sum(scheduled_minutes)::bigint AS scheduled_minutes, sum(inspected_units)::bigint AS inspected_units,
                  sum(rejected_units)::bigint AS rejected_units, sum(planned_downtime_minutes)::bigint AS planned,
                  sum(unplanned_downtime_minutes)::bigint AS unplanned, sum(failures)::bigint AS failures
           FROM v_equipment_daily
           WHERE organisation_id = %s AND day BETWEEN %s AND %s AND (%s::uuid IS NULL OR plant_id = %s)
           GROUP BY day ORDER BY day""", (organisation_id, start, end, plant_id, plant_id)).fetchall()
    reliability = [r for r in conn.execute("SELECT * FROM fn_equipment_reliability(%s, %s, %s)", (organisation_id, start, end)).fetchall()
                   if plant_id is None or r["plant_id"] == plant_id]
    reasons = conn.execute(
        """SELECT d.reason_code, d.category, count(*) AS events, sum(d.duration_minutes)::bigint AS minutes
           FROM downtime_events d JOIN equipment e ON e.id = d.equipment_id
           JOIN production_lines ln ON ln.id = e.production_line_id JOIN plants pl ON pl.id = ln.plant_id
           WHERE pl.organisation_id = %s AND (d.start_time AT TIME ZONE 'UTC')::date BETWEEN %s AND %s AND (%s::uuid IS NULL OR pl.id = %s)
           GROUP BY d.reason_code, d.category ORDER BY minutes DESC, d.reason_code""", (organisation_id, start, end, plant_id, plant_id)).fetchall()
    anomalies = detect(conn, organisation_id, start, end, "robust_z", plant_id=plant_id, limit=200) if include_anomalies else None

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    # ---- Daily (first, because Summary formulas point at it)
    ws = wb.create_sheet("Daily")
    _header(ws, 1, [c[0] for c in DAILY_COLUMNS])
    for r, row in enumerate(daily, 2):
        for c, (_, key, fmt) in enumerate(DAILY_COLUMNS, 1):
            cell = ws.cell(row=r, column=c, value=row[key])
            cell.number_format = fmt
    last = len(daily) + 1
    _widths(ws, [12] + [16] * 9)

    def total(col: str) -> str:
        return f"SUM(Daily!{col}2:{col}{max(last, 2)})"

    # ---- Summary
    summary["A1"] = "OpsMind operations report"
    summary["A1"].font = Font(bold=True, size=16)
    summary["A2"] = "Synthetic demonstration data unless you imported your own records."
    summary["A2"].font = Font(italic=True, color="9A3412")
    meta = [("Period", f"{start.isoformat()} to {end.isoformat()}"), ("Scope", plant_name or "All plants"),
            ("Days with production records", len(daily)), ("Generated (UTC)", generated_at.strftime("%Y-%m-%d %H:%M"))]
    for i, (k, v) in enumerate(meta, 4):
        summary.cell(row=i, column=1, value=k).font = Font(bold=True)
        summary.cell(row=i, column=2, value=v)
    summary["A9"], summary["B9"], summary["C9"] = "KPI", "Value", "Definition"
    for c in "ABC":
        summary[f"{c}9"].fill, summary[f"{c}9"].font = HEADER_FILL, HEADER_FONT
    U, T, R, S, I, J, P, D, F = (total(c) for c in "BCDEFGHIJ")
    kpis = [
        ("Units produced", f"={U}", "#,##0", "Sum of units produced"),
        ("Target attainment", f'=IF({T}>0,{U}/{T},"n/a")', "0.0%", "Units produced ÷ target units"),
        ("Utilisation", f'=IF({S}>0,{R}/{S},"n/a")', "0.0%", "Run minutes ÷ scheduled minutes"),
        ("Availability", f'=IF({S}-{P}>0,MAX(0,({S}-{P}-{D})/({S}-{P})),"n/a")', "0.0%",
         "(Scheduled − planned downtime − unplanned downtime) ÷ (scheduled − planned downtime), floored at 0"),
        ("Defect rate", f'=IF({I}>0,{J}/{I},"n/a")', "0.00%", "Rejected ÷ inspected units (only records with quality data)"),
        ("Unplanned downtime (min)", f"={D}", "#,##0", "Downtime attributed to the UTC date it started"),
        ("Unplanned failures", f"={F}", "#,##0", "Count of unplanned downtime events"),
    ]
    for i, (label, formula, fmt, definition) in enumerate(kpis, 10):
        summary.cell(row=i, column=1, value=label)
        cell = summary.cell(row=i, column=2, value=formula)
        cell.number_format = fmt
        cell.alignment = Alignment(horizontal="right")
        summary.cell(row=i, column=3, value=definition)
    _widths(summary, [30, 24, 90])

    # ---- Reliability
    rs = wb.create_sheet("Reliability")
    labels = ["Plant", "Line", "Asset", "Equipment", "Failures", "Unplanned (min)", "Planned (min)", "Run (min)", "Scheduled (min)",
              "MTBF (min)", "MTTR (min)", "Availability", "Rank in plant", "% of plant unplanned", "Cumulative %"]
    _header(rs, 1, labels)
    for r, x in enumerate(reliability, 2):
        vals = [x["plant_name"], x["line_name"], x["asset_code"], x["equipment_name"], x["failures"], x["unplanned_minutes"], x["planned_minutes"],
                x["run_minutes"], x["scheduled_minutes"], x["mtbf_minutes"], x["mttr_minutes"], x["availability"],
                x["downtime_rank_in_plant"], x["pct_of_plant_unplanned"], x["cumulative_pct_of_plant"]]
        for c, v in enumerate(vals, 1):
            cell = rs.cell(row=r, column=c, value=float(v) if hasattr(v, "as_tuple") else v)   # Decimal -> float
            if c == 12:
                cell.number_format = "0.0%"
    _widths(rs, [14, 12, 10, 26] + [14] * 11)
    rs.cell(row=len(reliability) + 3, column=1,
            value="MTBF is blank when a machine had no failure in the period (not infinity). Availability is blank when nothing was scheduled.")

    # ---- Downtime reasons
    ds = wb.create_sheet("Downtime reasons")
    _header(ds, 1, ["Reason code", "Category", "Events", "Minutes", "Share of listed minutes"])
    for r, x in enumerate(reasons, 2):
        ds.cell(row=r, column=1, value=x["reason_code"])
        ds.cell(row=r, column=2, value=x["category"])
        ds.cell(row=r, column=3, value=x["events"])
        ds.cell(row=r, column=4, value=x["minutes"])
        c = ds.cell(row=r, column=5, value=f'=IF(SUM($D$2:$D${len(reasons) + 1})>0,D{r}/SUM($D$2:$D${len(reasons) + 1}),"n/a")')
        c.number_format = "0.0%"
    _widths(ds, [20, 12, 10, 12, 22])

    # ---- Anomalies
    if anomalies is not None:
        an = wb.create_sheet("Anomalies")
        _header(an, 1, ["Date", "Asset", "Metric", "Value", "Baseline", "Score", "Explanation"])
        for r, x in enumerate(anomalies["anomalies"], 2):
            for c, v in enumerate([x["date"], x["assetCode"], x["metric"], x["value"], x["baseline"], x["score"], x["explanation"]], 1):
                an.cell(row=r, column=c, value=v)
        n = len(anomalies["anomalies"]) + 3
        an.cell(row=n, column=1, value=f"{anomalies['totalFlags']} flag(s) in range, {anomalies['returned']} listed. Method: robust_z. "
                                      "Flags are prompts to investigate, not diagnoses.")
        _widths(an, [12, 10, 16, 10, 10, 8, 110])

    # ---- About
    ab = wb.create_sheet("About")
    for i, line in enumerate([
        "How to read this workbook",
        "Summary KPIs are Excel formulas over the Daily sheet; click a cell to see the formula.",
        "Days without any production record do not appear in Daily (downtime on such days is not included — a documented limitation).",
        "Downtime is attributed to the UTC date on which the event starts; events crossing midnight are not split.",
        "Ratios are shown as n/a when the denominator is zero, never as 0%.",
        "Anomaly flags compare each machine with its own recent history (see docs/ANOMALY_DETECTION.md).",
    ], 1):
        ab.cell(row=i, column=1, value=line)
    ab["A1"].font = Font(bold=True)
    ab.column_dimensions["A"].width = 120

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
